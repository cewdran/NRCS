#!/usr/bin/env python3
"""Recompute derived input means and diversity on HPC from frozen source values."""
from hpc_guard import require_hpc
execution=require_hpc()
from pathlib import Path
import json,hashlib,datetime
from openpyxl import load_workbook
from pypdf import PdfReader
from source_identity import validate_mean_sources
import numpy as np
import pandas as pd

base=Path(__file__).resolve().parent
path=base/'inputs/workbook_data.json'; original=base/'inputs/workbook_data_supplied.json'
workbook=base/'inputs/NRCS_All_Available_Data_2026-09-29.xlsx'
if path.exists() or original.exists():
 raise SystemExit('Prepared input already exists. Submit a new isolated run; do not overwrite provenance.')
# Extract the canonical workbook only within the Slurm allocation. Cached source
# means are retained for comparison; all derived means and alpha are recomputed below.
def clean(value):
 if isinstance(value,(datetime.datetime,datetime.date)):return value.strftime('%Y-%m-%d')
 return value
wb=load_workbook(workbook,read_only=True,data_only=True)
t={}
for ws in wb.worksheets:
 if ws.title=='Read me':continue
 rows=ws.iter_rows(values_only=True); columns=[str(v) for v in next(rows)]
 assert len(columns)==len(set(columns)),ws.title
 t[ws.title]={'columns':columns,'rows':[[clean(v) for v in row] for row in rows if any(v is not None for v in row)]}
wb.close()
data={'source':'Canonical corrected workbook, version 2; extracted on HPC','tables':t}
original.write_text(json.dumps(data,ensure_ascii=False,allow_nan=False)+'\n')
def frame(name):return pd.DataFrame(t[name]['rows'],columns=t[name]['columns'])
r=frame('Sample records').set_index('record_id');m=frame('Measurements')
assert len(r)==303 and len(m)==6204
ward=r[(r.preferred_pdf_source_id=='S33') & (r.dataset=='Whipstone PODS soil')]
assert len(ward)==28 and ward.confirmed_sample_key.nunique()==28
assert (ward.collection_date=='2023-11-16').all()
assert r.loc['R0155','confirmed_sample_key']=='Blackberry_2B|2023-11-16'
assert r.loc['R0156','confirmed_sample_key']=='Blackberry_2C|2023-11-16'
assert r.loc['R0232','candidate_sample_key']=='Blackberry_2D|2025-08-17'
assert r.loc['R0232','original_lab_id']=='95559' or str(r.loc['R0232','original_lab_id'])=='95559'
assert r.loc['R0134','candidate_sample_key']=='GR_R7_T8|2024-11-06'
assert r.loc['R0134','collection_date'] is None or pd.isna(r.loc['R0134','collection_date'])
assert r.loc['R0073','pool_member_count']==6 and r.loc['R0074','pool_member_count']==6
ward_fields=m[(m.preferred_pdf_source_id=='S33') & m.value_basis.fillna('').str.startswith('PDF')]
assert len(ward_fields)==780 and ward_fields.record_id.nunique()==30
assert (ward_fields.groupby('record_id').size()==26).all()
pdf=PdfReader(base/'inputs/Whipstone_Soil_12.2023.pdf');assert len(pdf.pages)==30
for _,rec in r[r.preferred_pdf_source_id=='S33'].iterrows():
 text=pdf.pages[int(rec.preferred_pdf_page)-1].extract_text()
 assert str(rec.original_lab_id) in text,(rec.original_lab_id,rec.preferred_pdf_page)
(base.parent/'audit/corrected_records.json').write_text(json.dumps({
 'execution':execution,'workbook_sha256':hashlib.sha256(workbook.read_bytes()).hexdigest(),
 'sample_records':len(r),'measurement_rows':len(m),'Whipstone_2023_experimental_records':len(ward),
 'Ward_PDF_pages':len(pdf.pages),'Ward_reported_fields':len(ward_fields),
 'Whipstone_2023_collection_date':'2023-11-16',
 'date_scope':'Only researcher-confirmed Whipstone 2023 matches; candidate matches in other events are not promoted.'},indent=2)+'\n')
ms=t['Measurements'];mc={c:i for i,c in enumerate(ms['columns'])};means=t['Oatman soil means'];hc={c:i for i,c in enumerate(means['columns'])}
mean_rows=0;max_mean_delta=0.0;soil_only_identity=[]
source_records=r.to_dict('index')
mt=t['16S samples'];mi=mt['columns'].index('sample_key')
microbiome_keys={row[mi] for row in mt['rows']}
for row in means['rows']:
 indices=[int(x.strip())-2 for x in str(row[hc['measurement_excel_rows']]).split(',')]
 source=[ms['rows'][i]for i in indices]
 identity=validate_mean_sources(dict(zip(means['columns'],row)),
  [dict(zip(ms['columns'],v)) for v in source],source_records,microbiome_keys)
 if identity=='verified_soil_only_without_16S_match':
  soil_only_identity.append({'sample_key':row[hc['sample_key']],'variable':row[hc['variable']],
   'record_id':'R0025','laboratory_id':'34082','status':identity})
 values=np.array([r[mc['value']]for r in source],dtype=float)
 assert np.isfinite(values).all()and len(values)==int(row[hc['lab_record_n']])
 mean=float(values.mean());delta=abs(mean-float(row[hc['mean_value']]));max_mean_delta=max(max_mean_delta,delta)
 assert np.isclose(mean,row[hc['mean_value']],rtol=1e-10,atol=1e-10)
 row[hc['mean_value']]=mean;row[hc['sd_between_records']]=float(values.std(ddof=1))if len(values)>1 else None
 mean_rows+=1
meta=t['16S samples'];ix={c:i for i,c in enumerate(meta['columns'])};alpha_rows=0;max_shannon_delta=0.0
for domain in ['bacteria','archaea']:
 counts=pd.read_csv(base/'inputs'/f'{domain}_counts.tsv.gz',sep='\t',index_col=0)
 for row in meta['rows']:
  if row[ix['microbial_domain']]!=domain:continue
  sid=row[ix['analysis_sample']];expected=float(row[ix['target_read_count']])
  if sid not in counts.index:
   assert expected==0;continue
  c=counts.loc[sid].to_numpy(float);total=float(c.sum());assert total==expected
  richness=int((c>0).sum());assert richness==int(row[ix['observed_ASVs']])
  if total>0:
   p=c[c>0]/total;shannon=float(-np.sum(p*np.log(p)));prior=float(row[ix['Shannon']])
   max_shannon_delta=max(max_shannon_delta,abs(shannon-prior));assert np.isclose(shannon,prior,rtol=1e-10,atol=1e-10)
   row[ix['Shannon']]=shannon
  row[ix['target_read_count']]=total;row[ix['observed_ASVs']]=richness;alpha_rows+=1
path.write_text(json.dumps(data,ensure_ascii=False,allow_nan=False)+'\n')
audit={'execution':execution,'source_identity_check':'Laboratory record IDs, date, plot, variable and unit; explicit verified soil-only exception',
 'verified_soil_only_source_rows':soil_only_identity,'Haney_mean_rows_recomputed':mean_rows,'diversity_rows_recomputed':alpha_rows,
 'max_mean_difference_from_supplied':max_mean_delta,'max_Shannon_difference_from_supplied':max_shannon_delta,
 'supplied_sha256':hashlib.sha256(original.read_bytes()).hexdigest(),'recomputed_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
 'method':'Mean of explicitly indexed source measurements; Shannon=-sum(p*ln(p)) and observed positive-ASV richness from target counts.'}
(base.parent/'audit'/'derived_inputs.json').write_text(json.dumps(audit,indent=2)+'\n');print(json.dumps(audit,indent=2))
